#!/usr/bin/env python3
import argparse
import json
import logging
import math
import os
import re
import socket
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
import easyocr

reader = easyocr.Reader(['en'], gpu =True)

def parse_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    
    # 1. Handling File Gambar (.png, .jpg)
    if ext in ['.png', '.jpg', '.jpeg']:
        try:
            results = reader.readtext(file_path, detail=0)
            return " ".join(results)  # Mengembalikan teks hasil OCR dari gambar
        except Exception as e:
            print(f"[Warning] Gagal OCR gambar {file_path}: {e}")
            return ""
            
    # 2. Handling File Teks Biasa (.txt, .log, .py)
    elif ext in ['.txt', '.log', '.py']:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
            
    # ... (tambahkan parser pdf/docx/excel jika ada)
    return ""

SOCKET_PATH = "/tmp/rag_daemon.sock"
TCP_HOST = "127.0.0.1"
TCP_PORT = 8765
DEFAULT_MODEL = "sentence-transformers/all-MiniLM-L6-v2"

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")


@dataclass
class DocumentChunk:
    file_path: str
    text: str
    chunk_index: int = 0
    metadata: Dict[str, Any] = field(default_factory=dict)


class EmbeddingService:
    def __init__(self, model_name: str = DEFAULT_MODEL):
        self.model_name = model_name
        self.model = None
        self.device = "cuda" if self._cuda_available() else "cpu"
        try:
            from sentence_transformers import SentenceTransformer

            self.model = SentenceTransformer(model_name, device=self.device)
            logging.info("Loaded embedding model %s on %s", model_name, self.device)
        except Exception as exc:  # pragma: no cover - fallback path
            logging.warning("SentenceTransformers unavailable: %s", exc)
            self.model = None

    @staticmethod
    def _cuda_available() -> bool:
        try:
            import torch

            return torch.cuda.is_available()
        except Exception:
            return False

    def embed(self, texts: List[str]) -> List[List[float]]:
        if not texts:
            return []
        if self.model is not None:
            try:
                vectors = self.model.encode(texts, convert_to_numpy=True, normalize_embeddings=True)
                return [list(map(float, vector)) for vector in vectors]
            except Exception as exc:
                logging.warning("Embedding call failed; falling back to hash embeddings: %s", exc)
        return [self._hash_embedding(text) for text in texts]

    @staticmethod
    def _hash_embedding(text: str, dim: int = 384) -> List[float]:
        vec = [0.0 for _ in range(dim)]
        tokens = re.findall(r"[A-Za-z0-9]+", text.lower())
        if not tokens:
            return vec
        for token in tokens:
            idx = abs(hash(token)) % dim
            vec[idx] += 1.0
        norm = math.sqrt(sum(v * v for v in vec))
        if norm > 0:
            vec = [v / norm for v in vec]
        return vec


class VectorIndex:
    def __init__(self):
        self.chunks: List[DocumentChunk] = []
        self.embeddings: List[List[float]] = []

    def clear(self):
        self.chunks.clear()
        self.embeddings.clear()

    def add(self, chunk: DocumentChunk, embedding: List[float]):
        self.chunks.append(chunk)
        self.embeddings.append(embedding)

    def build_from_corpus(self, corpus_dir: str, embedding_service: EmbeddingService) -> None:
        self.clear()
        corpus_path = Path(corpus_dir)
        if not corpus_path.exists() or not corpus_path.is_dir():
            logging.warning("Corpus directory not found: %s", corpus_dir)
            return

        files = self._discover_files(corpus_path)
        texts_to_embed: List[Tuple[str, str, int, Dict[str, Any]]] = []
        for file_path in files:
            try:
                extracted = parse_file(file_path)
            except Exception as exc:  # pragma: no cover - file-specific resilience
                logging.warning("Failed to parse %s: %s", file_path, exc)
                continue
            if not extracted:
                continue
            for idx, chunk_text in enumerate(extracted):
                relative = os.path.relpath(file_path, corpus_path)
                texts_to_embed.append((str(relative), chunk_text, idx, {"source_file": file_path}))

        if not texts_to_embed:
            logging.info("No text chunks indexed from %s", corpus_dir)
            return

        batch = [chunk_text for _, chunk_text, _, _ in texts_to_embed]
        vectors = embedding_service.embed(batch)
        for (relative_path, chunk_text, chunk_idx, metadata), vector in zip(texts_to_embed, vectors):
            self.add(DocumentChunk(file_path=relative_path, text=chunk_text, chunk_index=chunk_idx, metadata=metadata), vector)
        logging.info("Indexed %d chunks across %d files in %s", len(self.chunks), len(files), corpus_dir)

    @staticmethod
    def _discover_files(root: Path) -> List[Path]:
        files: List[Path] = []
        if not root.exists():
            return files
        for current, _, filenames in os.walk(root):
            for name in sorted(filenames):
                full_path = Path(current) / name
                if full_path.is_file():
                    files.append(full_path)
        return files

    def query(self, query_text: str, top_k: int = 5, embedding_service: Optional[EmbeddingService] = None) -> List[Tuple[float, DocumentChunk]]:
        if not self.chunks:
            return []
        service = embedding_service or EmbeddingService()
        query_vector = service.embed([query_text])[0]
        scored: List[Tuple[float, DocumentChunk]] = []
        for chunk, vector in zip(self.chunks, self.embeddings):
            score = cosine_similarity(query_vector, vector)
            scored.append((score, chunk))
        scored.sort(key=lambda item: item[0], reverse=True)
        return scored[:top_k]


def cosine_similarity(vec_a: List[float], vec_b: List[float]) -> float:
    if not vec_a or not vec_b:
        return 0.0
    norm_a = math.sqrt(sum(v * v for v in vec_a))
    norm_b = math.sqrt(sum(v * v for v in vec_b))
    if norm_a == 0 or norm_b == 0:
        return 0.0
    dot = sum(a * b for a, b in zip(vec_a, vec_b))
    return dot / (norm_a * norm_b)


def normalize_answer(value: Any) -> str:
    if value is None:
        return ""
    answer = str(value).strip().upper()
    for ch in "-._·":
        answer = answer.replace(ch, "")
    return answer.strip()


def safe_read_utf8(path: Path) -> str:
    try:
        with path.open("r", encoding="utf-8", errors="strict") as fh:
            return fh.read()
    except Exception:
        try:
            with path.open("r", encoding="utf-8", errors="replace") as fh:
                return fh.read()
        except Exception:
            return ""


def split_chunks(text: str, chunk_size: int = 900) -> List[str]:
    if not text:
        return []
    paragraphs = re.split(r"\n\s*\n+|\r\n\s*\r\n+", text)
    chunks: List[str] = []
    current = ""
    for paragraph in paragraphs:
        candidate = paragraph.strip()
        if not candidate:
            continue
        if len(current) + len(candidate) + 1 <= chunk_size:
            if current:
                current += "\n"
            current += candidate
        else:
            if current:
                chunks.append(current)
            current = candidate
    if current:
        chunks.append(current)
    if not chunks:
        chunks = [text[:chunk_size]]
    return chunks


def parse_pdf(path: Path) -> List[str]:
    texts: List[str] = []
    try:
        import pypdf

        reader = pypdf.PdfReader(str(path))
        for page in reader.pages:
            try:
                page_text = page.extract_text() or ""
                if page_text:
                    texts.append(page_text)
            except Exception:
                continue
    except Exception:
        pass
    if texts:
        return split_chunks("\n\n".join(texts))

    try:
        import fitz

        doc = fitz.open(str(path))
        for page in doc:
            try:
                page_text = page.get_text("text")
                if page_text:
                    texts.append(page_text)
            except Exception:
                continue
        doc.close()
    except Exception:
        pass
    if texts:
        return split_chunks("\n\n".join(texts))

    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                try:
                    page_text = page.extract_text() or ""
                    if page_text:
                        texts.append(page_text)
                except Exception:
                    continue
    except Exception:
        pass
    if texts:
        return split_chunks("\n\n".join(texts))

    return []


def parse_docx(path: Path) -> List[str]:
    try:
        from docx import Document

        doc = Document(str(path))
        pieces: List[str] = []
        for element in doc.paragraphs:
            text = (element.text or "").strip()
            if text:
                pieces.append(text)
        for table in doc.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if values:
                    pieces.append(" | ".join(values))
        return split_chunks("\n".join(pieces))
    except Exception as exc:
        logging.warning("Failed to parse DOCX %s: %s", path, exc)
        return []


def parse_excel_or_csv(path: Path) -> List[str]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".csv":
            import pandas as pd

            frame = pd.read_csv(path, dtype=str, keep_default_na=False)
            rows = []
            for _, row in frame.iterrows():
                rows.append(" | ".join(str(cell).strip() for cell in row.tolist() if str(cell).strip()))
            return split_chunks("\n".join(rows))

        if suffix in {".xlsx", ".xls"}:
            import pandas as pd

            frame = pd.read_excel(path, engine="openpyxl")
            rows = []
            for _, row in frame.fillna("").iterrows():
                rows.append(" | ".join(str(cell).strip() for cell in row.tolist() if str(cell).strip()))
            return split_chunks("\n".join(rows))
        return []
    except Exception as exc:
        logging.warning("Failed to parse spreadsheet %s: %s", path, exc)
        return []


def parse_text_like(path: Path) -> List[str]:
    try:
        content = safe_read_utf8(path)
        return split_chunks(content)
    except Exception as exc:
        logging.warning("Failed to read text-like file %s: %s", path, exc)
        return []


def parse_image(path: Path) -> List[str]:
    try:
        import easyocr

        reader = easyocr.Reader(["en"], gpu=False)
        results = reader.readtext(str(path), detail=0, paragraph=True)
        if results:
            return split_chunks("\n".join(results))
    except Exception:
        pass

    try:
        from PIL import Image
        from transformers import pipeline

        pipe = pipeline("image-to-text", model="google/tapas-base-finetuned-wtq")
        image = Image.open(path)
        texts = pipe(image)
        if texts:
            values = [item.get("generated_text", "") for item in texts if isinstance(item, dict)]
            result = "\n".join(v for v in values if v)
            if result:
                return split_chunks(result)
    except Exception:
        pass

    return []


def parse_file(path: Path) -> List[str]:
    if not path.exists() or not path.is_file():
        return []
    suffix = path.suffix.lower()

    try:
        if suffix == ".pdf":
            return parse_pdf(path)
        if suffix == ".docx":
            return parse_docx(path)
        if suffix in {".xlsx", ".xls", ".csv"}:
            return parse_excel_or_csv(path)
        if suffix in {".txt", ".log", ".py"}:
            return parse_text_like(path)
        if suffix in {".png", ".jpg", ".jpeg"}:
            return parse_image(path)
    except Exception as exc:
        logging.warning("Unhandled parse error for %s: %s", path, exc)
        return []
    return []


def extract_question_keywords(question: str) -> List[str]:
    text = question.lower()
    return [token for token in re.findall(r"[a-z0-9]+", text) if len(token) > 2]


def candidate_lines_from_chunk(chunk_text: str) -> List[str]:
    lines = []
    for line in re.split(r"\n+|\r+", chunk_text):
        cleaned = line.strip()
        if not cleaned or len(cleaned) < 3:
            continue
        if any(ch.isdigit() for ch in cleaned) or re.search(r"\b[A-Z]{2,}[A-Z0-9_-]{2,}\b", cleaned):
            lines.append(cleaned)
    return lines


def is_likely_value(candidate: str) -> bool:
    value = candidate.strip().strip(".,;:()[]{}\"'")
    if not value:
        return False
    if value.upper() in {"PART", "NUMBER", "VALUE", "NAME", "MODEL", "REVISION", "VERSION", "TYPE", "CODE", "SERIAL", "ID"}:
        return False
    if re.search(r"\d", value):
        return True
    if re.fullmatch(r"[A-Z0-9][A-Z0-9._/-]*", value) and len(value) >= 2:
        return True
    return False


def extract_value_candidates(line: str) -> List[str]:
    candidates: List[str] = []
    normalized = line.strip().rstrip(".,;:)]}\"'")
    for sep in (":", "="):
        if sep in normalized:
            right = normalized.split(sep, 1)[1].strip()
            for part in re.split(r"\s*\|\s*|\s*;\s*|\s*,\s*", right):
                candidate = part.strip().strip(".,;:()[]{}\"'")
                if is_likely_value(candidate):
                    candidates.append(candidate)
            if candidates:
                return candidates

    for pattern in [
        r"\b(?:Q\d+|FY\d{2,4}|REV-[A-Z0-9]+|ORR-[A-Z0-9-]+|[A-Z]{2,}[A-Z0-9_-]{2,})\b",
        r"\b\d+(?:\.\d+)?(?:%|V|A|W|Hz|KHz|MHz|GHz|mA|mV|ms|s|C|MB|GB|TB|KB|B|RPM|mm|nm)?\b",
    ]:
        for match in re.finditer(pattern, normalized, flags=re.IGNORECASE):
            candidate = match.group(0).strip()
            if is_likely_value(candidate):
                candidates.append(candidate)
    return candidates


def extract_answer_from_context(question: str, matches: List[DocumentChunk]) -> str:
    if not matches:
        return ""
    keywords = set(extract_question_keywords(question))
    best_line = ""
    best_score = -1
    best_candidates: List[str] = []

    for match in matches:
        for line in candidate_lines_from_chunk(match.text):
            line_lower = line.lower()
            score = 0
            for keyword in keywords:
                if keyword in line_lower:
                    score += 10
            if any(ch.isdigit() for ch in line):
                score += 4
            if re.search(r"\b[A-Z]{2,}[A-Z0-9_-]{2,}\b", line):
                score += 2
            values = extract_value_candidates(line)
            if values:
                score += 6
            if score > best_score:
                best_score = score
                best_line = line
                best_candidates = values

    if not best_line:
        first_line = candidate_lines_from_chunk(matches[0].text)[0] if candidate_lines_from_chunk(matches[0].text) else ""
        if first_line:
            best_line = first_line
            best_candidates = extract_value_candidates(first_line)

    if best_candidates:
        return normalize_answer(best_candidates[0])

    value_match = re.search(r"(?<![A-Z0-9])(?:\d+(?:\.\d+)?(?:\s*(?:%|MS|MB|GB|W|V|A|HZ|KHZ|MHZ|GHZ|C|MA|MV|DB|RPM|NM|S|H|D|N|KG|MM|IN|LB|KW|TB|KB|B))?|[A-Z]{2,}[A-Z0-9_-]{2,})", best_line, flags=re.IGNORECASE)
    if value_match:
        return normalize_answer(value_match.group(0))
    return ""


def build_response_for_query(index: VectorIndex, question: str, embedding_service: Optional[EmbeddingService] = None) -> Dict[str, Any]:
    matches = index.query(question, top_k=5, embedding_service=embedding_service)
    if not matches:
        return {"answer": "", "citations": [], "confidence": 0.0}

    top_chunks = [chunk for _, chunk in matches]
    answer = extract_answer_from_context(question, top_chunks)
    citations = sorted({chunk.file_path for chunk in top_chunks if chunk.file_path})
    if not answer:
        return {"answer": "", "citations": [], "confidence": 0.0}
    confidence = max(0.1, min(0.99, 0.35 + (matches[0][0] if matches else 0.0) * 0.65))
    return {"answer": answer, "citations": citations, "confidence": round(confidence, 4)}


class RAGDaemon:
    def __init__(self):
        self.embedding_service = EmbeddingService()
        self.index = VectorIndex()
        self.index_path: Optional[str] = None

    def handle_index(self, corpus_dir: str) -> Dict[str, Any]:
        start = time.time()
        self.index.build_from_corpus(corpus_dir, self.embedding_service)
        self.index_path = corpus_dir
        elapsed = time.time() - start
        logging.info("Index built in %.2fs for %s", elapsed, corpus_dir)
        return {"status": "ok", "indexed": len(self.index.chunks), "elapsed_seconds": round(elapsed, 2)}

    def handle_query(self, corpus_dir: str, query_text: str) -> Dict[str, Any]:
        if not self.index.chunks or corpus_dir != self.index_path:
            self.index.build_from_corpus(corpus_dir, self.embedding_service)
        response = build_response_for_query(self.index, query_text, embedding_service=self.embedding_service)
        return response

    def serve(self, socket_path: str = SOCKET_PATH):
        if hasattr(socket, "AF_UNIX"):
            try:
                os.unlink(socket_path)
            except FileNotFoundError:
                pass
            server = socket.socket(socket.AF_UNIX, socket.SOCK_STREAM)
            server.bind(socket_path)
            address = socket_path
        else:
            server = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            server.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            server.bind((TCP_HOST, TCP_PORT))
            address = f"{TCP_HOST}:{TCP_PORT}"

        with server:
            server.listen(32)
            logging.info("RAG daemon listening on %s", address)
            while True:
                try:
                    conn, _ = server.accept()
                    raw = b""
                    while True:
                        chunk = conn.recv(65536)
                        if not chunk:
                            break
                        raw += chunk
                        if len(chunk) < 65536:
                            break
                    try:
                        payload = json.loads(raw.decode("utf-8"))
                    except Exception:
                        payload = {"cmd": "noop"}
                    cmd = payload.get("cmd")
                    if cmd == "index":
                        response = self.handle_index(payload.get("corpus", ""))
                    elif cmd == "query":
                        response = self.handle_query(payload.get("corpus", ""), payload.get("query", ""))
                    else:
                        response = {"answer": "", "citations": [], "confidence": 0.0}
                    conn.sendall(json.dumps(response, ensure_ascii=False).encode("utf-8"))
                    conn.close()
                except Exception as exc:
                    logging.warning("Daemon socket loop exception: %s", exc)


def main() -> int:
    parser = argparse.ArgumentParser(description="RAG daemon for multimodal document indexing and retrieval.")
    parser.add_argument("--serve", action="store_true", help="Run the Unix socket daemon.")
    parser.add_argument("--index", type=str, help="Index a corpus directory and exit.")
    parser.add_argument("--corpus", type=str, help="Corpus directory for one-off queries.")
    parser.add_argument("--query", type=str, help="Query text to answer in one-off mode.")
    args = parser.parse_args()

    daemon = RAGDaemon()

    if args.serve:
        daemon.serve()
        return 0

    if args.index:
        response = daemon.handle_index(args.index)
        print(json.dumps(response, ensure_ascii=False))
        return 0

    if args.corpus and args.query:
        response = daemon.handle_query(args.corpus, args.query)
        print(json.dumps(response, ensure_ascii=False))
        return 0

    parser.print_help()
    return 1


if __name__ == "__main__":
    sys.exit(main())
